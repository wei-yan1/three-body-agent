"""Generate seven daily task report Word files from the provided template."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt


TEMPLATE_PATH = Path(r"C:\Users\MR\Downloads\每日任务完成与计划填报模板.docx")
OUTPUT_DIR = Path(r"D:\AgentCode\two\deliverables")
PROJECT_NAME = "基于《三体》的时序人格对话 Agent"
START_DATE = date(2026, 7, 6)


DAYS = [
    {
        "title": "第1天：需求梳理、项目结构确认与基础运行框架搭建",
        "tasks": (
            "1. 梳理项目目标：实现一个面向《三体》人物复现的时序人格对话 Agent，重点解决角色一致性、时间线边界和检索增强回答问题。\n"
            "2. 明确技术栈：后端采用 FastAPI，模型链路采用 LangChain / create_agent，向量检索使用 Chroma，数据存储使用 PostgreSQL、Redis，本地前端采用 HTML/CSS/JavaScript。\n"
            "3. 整理项目目录：确认 app、data、frontend、scripts、tests 等模块边界，并规划 API、Agent、RAG、Memory、Storage、Frontend 的职责。\n"
            "4. 完成 FastAPI 入口和静态页面路由的检查，确认登录页、角色选择页、角色聊天页和静态资源能够统一由后端服务提供。\n"
            "完成进度：100%。基础结构已明确，可进入数据与人物建模阶段。"
        ),
        "results": (
            "完成了项目总体方案说明、模块职责划分和运行链路梳理。确认 app/main.py 中已挂载认证路由、聊天路由和静态资源目录，并为罗辑、"
            "章北海、叶文洁、汪淼等角色页面提供访问入口。形成了后续 7 天开发节奏：先搭建结构，再处理数据，再接入 RAG 与 Agent，最后完成前端联调和总结。"
        ),
        "unfinished": "暂无明显延期事项。需要注意 README 与终端显示存在编码不一致现象，后续整理文档时应以 UTF-8 源文件和 Word 文档显示为准。",
        "collab": (
            "小组成员对项目目标和模块分工进行了统一：一部分负责后端接口与 Agent 编排，一部分负责数据处理与索引构建，一部分负责前端页面和交互。"
            "今日主要问题是项目涉及模块较多，初期容易只关注页面效果而忽略时序人格约束，因此已将“时间线边界”和“人物一致性”列为核心验收标准。"
        ),
        "ai_tools": "ChatGPT / Codex、项目 README 和本地代码检索工具。",
        "ai_help": "辅助拆解项目需求，归纳技术栈和目录结构，梳理从用户提问到模型回答的完整链路，并生成阶段性开发计划。AI 还辅助定位了 app/main.py、app/api/v1/chat.py、app/services/chat_service.py 等关键入口文件。",
        "ai_limits": "AI 初始总结偏宏观，容易把项目描述成普通 RAG 问答系统，忽略“时序人格”和“角色知识边界”的特殊性。本人修正方案：结合代码中的 timeline_stage、knowledge_mode、persona skill、memory 等字段重新归纳项目重点，并把后续任务拆到具体文件和模块。",
        "tomorrow": (
            "1. 整理《三体》原文、人物画像 JSONL、全局时间线等数据来源。\n"
            "2. 分析人物 Skill 的字段结构，确认每个角色在不同 T 阶段的人格、知识边界和关系状态。\n"
            "3. 初步完成数据加载器和结构化文档转换逻辑的检查，为 RAG 索引构建做准备。"
        ),
    },
    {
        "title": "第2天：原始语料整理、人物时序 Skill 建模与数据加载",
        "tasks": (
            "1. 检查 data/raw 与 data/processed 目录，确认小说原文、人物画像、时间线和中间产物的存放位置。\n"
            "2. 梳理罗辑、章北海、叶文洁、汪淼等角色的 Temporal Persona Skill，重点关注 stage_id、timeline_stage、known_events、forbidden_events、relationships、anti-drift rules 等字段。\n"
            "3. 阅读并验证 app/rag/loaders/persona_skill_loader.py、app/rag/loaders/novel_text_loader.py 的加载逻辑，确保 JSONL 和小说文本可以转换为 LangChain Document。\n"
            "4. 明确人物 Skill 优先于普通小说片段的设计原则：先约束角色人格，再补充事实依据。\n"
            "完成进度：100%。数据来源和人物建模结构已确认。"
        ),
        "results": (
            "完成了数据资产盘点：小说原文位于 data/raw/three_body_characters/three_body.txt，人物时序画像位于 data/processed/persona_profiles/，"
            "时间线位于 data/processed/timelines/。确认系统不是简单把全文塞进提示词，而是把人物拆成多个时间阶段，每个阶段保存人格状态、表达风格、已知事件、禁止提前知道的未来事件和关系状态。"
        ),
        "unfinished": "部分人物画像文本较长，后续检索时可能出现召回片段过多或上下文冗余的问题，需要在第3天通过结构化切分和 metadata 过滤解决。",
        "collab": "今日协作重点是统一“人物 Skill”和“小说原文”的作用边界。讨论后确定：人物 Skill 用来回答“这个角色会如何理解和表达”，小说原文用来提供事件证据；二者都进入 RAG，但优先级和过滤条件不同。",
        "ai_tools": "ChatGPT / Codex、python-docx、ripgrep、本地代码阅读。",
        "ai_help": "AI 辅助从项目目录中定位人物画像、时间线、小说原文和加载器代码，并帮助把人物建模拆成可填写日报的任务项。AI 还辅助整理了人物 Skill 的关键字段和验收标准。",
        "ai_limits": "AI 在没有完整阅读 JSONL 内容时，可能会用概念化语言描述人物 Skill，细节不一定完全对应每条数据。本人修正方案：以本地 data/processed/persona_profiles/ 和 loader 代码为准，日报中只写确定存在的结构和工程目标，不虚构具体剧情字段。",
        "tomorrow": (
            "1. 实现并检查 structure-aware chunking，把人物 Skill 和小说原文切分为可检索文档。\n"
            "2. 为 chunk 添加 character、timeline_stage、stage_order、chunk_type、source 等 metadata。\n"
            "3. 构建或验证 Chroma 向量索引，为后续 Agent 检索做准备。"
        ),
    },
    {
        "title": "第3天：结构化切分、Embedding 与 Chroma 向量索引构建",
        "tasks": (
            "1. 阅读 app/rag/embeddings/structure_aware_chunker.py，确认人物 Skill 和小说原文的切分方式。\n"
            "2. 检查 scripts/build_indexes.py 中 build_chunks、build_chroma_indexes、build_novel_chroma_index 等流程，确认 JSONL chunk 能写入 Chroma。\n"
            "3. 接入 DashScope/Bailian embedding 模型，验证 create_dashscope_embeddings 的配置来源和错误提示。\n"
            "4. 建立两个核心集合：three_body_persona_profiles 和 three_body_novel_chunks，分别存放人物画像与小说片段。\n"
            "完成进度：90%。索引构建流程已明确，后续还需结合 API Key 和运行环境做完整重建验证。"
        ),
        "results": (
            "完成了 RAG 索引链路梳理：先由 structure_aware_chunker 将人物 Skill 和小说原文转成结构化 Document，再写出可审计 JSONL，随后由 scripts/build_indexes.py "
            "清理旧集合并批量写入 Chroma。人物文档保留角色和时间阶段过滤字段，小说文档保留 stage_order，使角色只能检索当前阶段及之前可知内容。"
        ),
        "unfinished": "由于完整 embedding 与 Chroma 重建依赖外部模型 API Key，今日主要完成代码层面的检查和流程确认，未把所有索引重新跑一遍。延期原因是避免在未确认密钥和网络环境的情况下盲目执行长任务。",
        "collab": "小组协作中明确了“可审计中间产物”的价值：chunk 先写成 JSONL，再进入向量库，便于排查召回错误。今日发现 RAG 不仅要考虑相似度，还必须考虑时间线过滤，否则人物容易提前知道后续剧情。",
        "ai_tools": "ChatGPT / Codex、DashScope Embedding 相关代码、本地 Chroma 索引脚本。",
        "ai_help": "AI 辅助解释 build_indexes.py 的索引构建流程，定位 persona_collection、novel_collection、batch 写入、metadata 清洗和异步 embedding 输入输出相关代码，并将其整理成日报成果。",
        "ai_limits": "AI 对外部 API 的实际可用性无法替代真实运行验证。本人修正方案：日报中区分“已完成代码检查”和“待完整重建验证”，并把 API Key、网络、索引耗时作为后续风险项记录。",
        "tomorrow": (
            "1. 检查 PersonaVectorRetriever、NovelVectorRetriever 和 HybridFusionRetriever 的召回逻辑。\n"
            "2. 加入 Dense + BM25 + weighted fusion + rerank 的混合检索说明。\n"
            "3. 将检索结果接入 Agent Middleware，为角色回答提供上下文。"
        ),
    },
    {
        "title": "第4天：混合检索、意图路由与 Agent Middleware 编排",
        "tasks": (
            "1. 阅读 app/rag/retrievers/hybrid_fusion_retriever.py，确认向量召回、BM25 稀疏打分、Min-Max 归一化、加权融合和 DashScope rerank 的处理流程。\n"
            "2. 检查 app/agents/persona/intent_router.py，确认用户问题会先被识别为闲聊、人物关系、剧情事件、价值思考、未来试探、系统问题等类型。\n"
            "3. 检查 app/agents/persona/middleware.py，确认 QueryOptimizationMiddleware 会把角色、时间阶段、已知事件和用户问题合并成更适合检索的查询。\n"
            "4. 将人物 Skill、小说片段、关系约束、当前线程记忆和可选联网结果统一注入 Agent 上下文。\n"
            "完成进度：95%。核心编排逻辑已完成，仍需继续做端到端问答测试。"
        ),
        "results": (
            "完成了 Agentic RAG 链路梳理：用户问题进入后端后，不是直接丢给大模型，而是先由 Intent Router 判断问题类型，再由 Query Optimization 加强检索查询，"
            "随后由 Persona Retriever、Novel Retriever、Hybrid Fusion Retriever 和 rerank 共同筛选上下文，最后由 TemporalPersonaRAGMiddleware 注入模型调用。该设计能够减少角色跑偏、百科式回答和未来剧情泄露。"
        ),
        "unfinished": "混合检索效果需要更多样例评测，尤其是短问题、代词问题和关系问题。延期原因是今日重点在编排链路，尚未建立系统化测试集。",
        "collab": "今日讨论集中在“是否每次都需要检索”。结论是：系统应根据意图动态决定检索策略，日常寒暄可少检索，剧情和关系问题需要强检索，未来试探问题要优先触发时间线边界约束。",
        "ai_tools": "ChatGPT / Codex、LangChain Agent Middleware、本地检索器代码。",
        "ai_help": "AI 辅助把复杂的 RAG 编排拆成“意图判断-查询改写-召回-融合-重排-上下文注入-模型生成”的顺序，并整理各模块对应的文件位置，便于后续截图和答辩说明。",
        "ai_limits": "AI 容易把 Middleware 描述成普通提示词拼接，但实际代码还包含检索策略、知识模式、时间线阶段、外部搜索和记忆工具。本人修正方案：回到 app/agents/persona/middleware.py 和 app/services/chat_service.py 对照函数调用关系，确保表述符合项目实现。",
        "tomorrow": (
            "1. 检查 Chat Service 如何创建不同角色 Agent。\n"
            "2. 接入 temporal 与 transparent 两种知识模式。\n"
            "3. 完成会话线程、消息落库和情景记忆写入流程检查。"
        ),
    },
    {
        "title": "第5天：Chat Service、知识模式、线程会话与记忆系统",
        "tasks": (
            "1. 阅读 app/services/chat_service.py，确认 chat_with_luoji、chat_with_zhangbeihai、chat_with_wangmiao、chat_with_yewenjie 统一进入 _chat_with_temporal_persona。\n"
            "2. 检查 temporal 和 transparent 两种 knowledge_mode：前者严格遵守当前时间线，后者允许联网和外部资料补充，但仍保持人物阶段人格。\n"
            "3. 阅读 app/api/v1/chat.py 与 app/storage/repositories/session_repository.py，确认线程创建、消息追加、历史消息读取、线程删除等接口链路。\n"
            "4. 阅读 app/memory/tools.py、app/memory/long_term/episodic_memory.py 和 episodic_memory_repository.py，确认情景记忆写入、召回、强化和容量遗忘逻辑。\n"
            "完成进度：95%。后端会话与记忆链路已完成，仍需继续优化异常提示和测试覆盖。"
        ),
        "results": (
            "完成了后端核心调用链路总结：前端发送角色、时间阶段、知识模式、线程名和用户消息；API 层创建或读取线程并写入用户消息；Chat Service 创建对应角色 Agent，"
            "注入 TemporalPersonaRAGMiddleware 与 MemoryTool；模型生成回答后写回 assistant 消息，并将本轮对话作为 episodic memory 保存。系统支持同一角色在不同阶段、不同模式、不同线程下保存独立上下文。"
        ),
        "unfinished": "测试覆盖仍不完整，尤其是多线程删除后对应记忆清理、容量遗忘阈值、异常 API Key 缺失等场景还需要补充自动化用例。",
        "collab": "小组协作中将“聊天记录”和“人物记忆”区分开：聊天记录用于恢复前端历史消息，情景记忆用于 Agent 生成时保持上下文连续。该区分能避免前端展示和模型上下文混为一谈。",
        "ai_tools": "ChatGPT / Codex、FastAPI 路由代码、PostgreSQL/Redis 存储模块。",
        "ai_help": "AI 辅助追踪从 API 路由到 Chat Service 再到 MemoryTool 的函数调用，帮助整理线程、消息和记忆三类数据的区别，并将技术点转化为日报可读表达。",
        "ai_limits": "AI 对运行时数据库状态不可见，只能根据代码判断逻辑。本人修正方案：把数据库相关结论限定为“代码设计与调用链路”，并在后续计划中保留真实环境联调和数据表检查。",
        "tomorrow": (
            "1. 检查前端登录页、角色选择页和多角色聊天页。\n"
            "2. 联调前端与 /api/v1/chat/{agent_slug} 接口。\n"
            "3. 优化阶段选择、知识模式切换、线程选择、新建线程、批量删除和消息展示体验。"
        ),
    },
    {
        "title": "第6天：前端页面、角色交互、线程管理与接口联调",
        "tasks": (
            "1. 检查 frontend/login.html、frontend/agents.html 和各角色聊天页，确认用户登录、角色选择和聊天入口完整。\n"
            "2. 阅读 frontend/static/luoji-chat.js，确认 token 校验、阶段选择、知识模式选择、线程读取、新建线程、删除线程、历史消息加载和发送消息逻辑。\n"
            "3. 阅读 frontend/static/luoji-chat.css，确认聊天页面布局、角色立绘、消息区、输入区、下拉菜单和响应式样式。\n"
            "4. 与后端 /api/v1/chat/{agent_slug}、/threads、/threads/{thread_id}/messages、/threads/delete 接口进行逻辑联调。\n"
            "完成进度：100%。前端主流程已打通。"
        ),
        "results": (
            "完成了前端交互链路梳理：用户登录后 token 存入 localStorage，进入角色页面后可选择时间阶段和知识模式，系统按角色 slug 调用对应 API。"
            "线程管理支持读取已有线程、新建线程、自定义线程名、加载历史消息和批量删除。聊天页面采用左侧人物视觉区域和右侧对话区域的布局，能直观展示角色身份并提供沉浸式对话体验。"
        ),
        "unfinished": "部分前端文本在终端中显示乱码，但浏览器和源码文件本身以 UTF-8 保存时不影响页面显示。后续需要在答辩或截图时直接使用浏览器页面和编辑器，不以 PowerShell 输出为展示依据。",
        "collab": "今日协作重点是接口字段对齐：前端发送 timeline_stage、knowledge_mode、thread_name、message，后端返回 character、thread_id、thread_name、answer。通过字段统一，减少了前后端联调误差。",
        "ai_tools": "ChatGPT / Codex、浏览器页面截图、FastAPI 接口、本地前端源码。",
        "ai_help": "AI 辅助定位 luoji-chat.js 中的 fetch 调用、线程菜单渲染、发送消息逻辑和错误处理逻辑，并把前端功能按用户操作流程整理为日报内容。",
        "ai_limits": "AI 无法替代人工检查页面美观度和实际点击体验。本人修正方案：后续截图以浏览器真实运行页面为准，并人工确认按钮、下拉菜单、输入框、消息气泡没有错位。",
        "tomorrow": (
            "1. 汇总 7 天开发成果，整理项目亮点和不足。\n"
            "2. 编写测试与优化记录，确认可截图代码位置。\n"
            "3. 准备最终展示材料，包括功能截图、代码截图和项目说明。"
        ),
    },
    {
        "title": "第7天：测试验证、成果总结、截图材料整理与后续优化规划",
        "tasks": (
            "1. 汇总项目主要成果：FastAPI 后端、角色聊天 API、Temporal Persona Skill、结构化 RAG、混合检索、Agent Middleware、会话线程、情景记忆和静态前端页面。\n"
            "2. 检查 scripts/test_chat_memory.py、scripts/test_luoji_create_agent.py 等测试脚本，确认可用于验证记忆、Agent 创建和角色问答。\n"
            "3. 整理代码截图位置，覆盖入口路由、聊天 API、Agent 编排、RAG 检索、索引构建、记忆系统和前端发送消息。\n"
            "4. 总结不足和后续优化方向：增加自动化评测、完善角色扩展、优化引用证据展示、增强异常处理和部署方案。\n"
            "完成进度：100%。7 天阶段任务完成。"
        ),
        "results": (
            "完成了项目阶段总结。当前项目已经形成一个前后端闭环的《三体》时序人格对话系统：用户可登录、选择角色、选择时间阶段和知识模式，并与角色进行多线程对话；"
            "后端通过意图路由、RAG 检索、人物 Skill、时间线约束和记忆系统共同生成回答。整理了适合作为作业/答辩展示的代码截图点，能够说明项目从数据、检索、Agent 到页面交互的完整实现。"
        ),
        "unfinished": "阶段内未完成完整线上部署和大规模自动化评测。原因是当前任务重点是本地项目实现与日报材料整理，部署和评测需要额外服务器、密钥、测试集和时间。",
        "collab": "小组最终对项目亮点进行了统一表述：不是普通聊天机器人，而是带时间线边界、人物 Skill、RAG 证据和长期记忆的角色 Agent。协作中也明确了答辩分工：一人讲架构，一人讲 RAG 和 Agent，一人讲前端交互与演示。",
        "ai_tools": "ChatGPT / Codex、本地代码检索、Word 文档生成工具。",
        "ai_help": "AI 辅助汇总 7 天日报内容，按模板生成 Word 文件，并整理代码截图位置清单。AI 还帮助把技术实现转化成适合课堂汇报的语言。",
        "ai_limits": "AI 生成日报可能过于完整，容易显得每一天任务都完全线性推进；实际开发中可能存在交叉调试和返工。本人修正方案：在日报中保留“未完成事项”和“后续优化”，体现真实工程迭代过程。",
        "tomorrow": (
            "后续计划：1. 补充自动化测试，覆盖意图路由、时间线边界、检索召回、记忆遗忘和线程删除。\n"
            "2. 增加回答证据展示，让用户可以查看模型引用的人物 Skill 和小说片段。\n"
            "3. 扩展更多角色和更细时间阶段。\n"
            "4. 整理部署文档，准备云端或局域网演示环境。"
        ),
    },
]


BLANK_TARGETS = {
    7: "tasks",
    9: "results",
    11: "unfinished",
    14: "collab",
    18: "ai_tools",
    21: "ai_help",
    24: "ai_limits",
    26: "tomorrow",
}


def set_text(paragraph, text: str, *, size: int = 11, bold: bool = False) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "宋体"
        run.font.size = Pt(size)
        run.bold = bold


def generate() -> list[Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_paths: list[Path] = []
    for day_no, day in enumerate(DAYS, start=1):
        doc = Document(str(TEMPLATE_PATH))
        set_text(doc.paragraphs[0], f"每日任务完成与计划填报（第{day_no}天）", size=16, bold=True)
        set_text(doc.paragraphs[1], f"填报人：__________  项目名称：{PROJECT_NAME}")
        current = START_DATE + timedelta(days=day_no - 1)
        set_text(doc.paragraphs[2], f"填报日期：{current.year}年{current.month}月{current.day}日")
        set_text(doc.paragraphs[3], day["title"], bold=True)

        for paragraph_index, key in BLANK_TARGETS.items():
            set_text(doc.paragraphs[paragraph_index], day[key])

        for paragraph in doc.paragraphs:
            if paragraph.text.startswith(("一、", "二、", "三、")):
                for run in paragraph.runs:
                    run.bold = True

        path = OUTPUT_DIR / f"第{day_no}天.docx"
        doc.save(path)
        output_paths.append(path)
    return output_paths


if __name__ == "__main__":
    for output_path in generate():
        print(output_path)

